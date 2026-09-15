import base64
import io
import qrcode
from flask import send_file
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from flask import (
    Blueprint,
    render_template,
    request,
    abort,
    send_file,
    redirect,
    url_for
)

from config import get_sql_server_connection


view_student_listing = Blueprint(
    "view_student_listing",
    __name__
)


# ============================================================
# HELPER: Convert database photo to data URL
# ============================================================

def photo_to_data_url(photo_binary):
    """
    Converts SQL Server VARBINARY(MAX) photo into
    browser-compatible base64 data URL.
    """

    if not photo_binary:
        return None

    try:
        encoded = base64.b64encode(photo_binary).decode("utf-8")

        return f"data:image/jpeg;base64,{encoded}"

    except Exception:
        return None


# ============================================================
# HELPER: Generate QR code from QrToken
# ============================================================

def generate_qr_image(qr_token):
    """
    Generates QR image using the SAME QrToken stored
    in SQL Server.
    """

    if not qr_token:
        return None

    qr = qrcode.QRCode(
        version=1,
        error_correction=qrcode.constants.ERROR_CORRECT_H,
        box_size=10,
        border=4
    )

    qr.add_data(str(qr_token))
    qr.make(fit=True)

    img = qr.make_image()

    output = io.BytesIO()

    img.save(output, format="PNG")

    output.seek(0)

    return output


# ============================================================
# STUDENT LISTING
# ============================================================

@view_student_listing.route(
    "/view_students",
    methods=["GET"]
)
def students():

    search = request.args.get(
        "search",
        ""
    ).strip()

    try:

        with get_sql_server_connection() as connection:

            cursor = connection.cursor()

            if search:

                sql = """
                SELECT
                    ID,
                    AdmissionNo,
                    FullName,
                    DOB,
                    Gender,
                    Mobile,
                    Email,
                    ProgramCode,
                    ProgramName,
                    PreferredBatch,
                    StudyMode,
                    CourseCompletionStatus,
                    QrToken,
                    PhotoBase64
                FROM tbl_student_admision
                WHERE
                    AdmissionNo LIKE ?
                    OR FullName LIKE ?
                    OR Mobile LIKE ?
                    OR Email LIKE ?
                ORDER BY AdmissionNo DESC
                """

                search_value = f"%{search}%"

                cursor.execute(
                    sql,
                    (
                        search_value,
                        search_value,
                        search_value,
                        search_value
                    )
                )

            else:

                sql = """
                SELECT
                    ID,
                    AdmissionNo,
                    FullName,
                    DOB,
                    Gender,
                    Mobile,
                    Email,
                    ProgramCode,
                    ProgramName,
                    PreferredBatch,
                    StudyMode,
                    CourseCompletionStatus,
                    QrToken,
                    PhotoBase64
                    
                FROM tbl_student_admision where CourseCompletionStatus not in('Hold','Completed','DROPOUT')
                ORDER BY AdmissionNo DESC
                """

                cursor.execute(sql)

            rows = cursor.fetchall()

            students_data = []

            for row in rows:

                student = {
                    "EnrollmentNo": row.ID,
                    "AdmissionNo": row.AdmissionNo,
                    "FullName": row.FullName,
                    "DOB": row.DOB,
                    "Gender": row.Gender,
                    "Mobile": row.Mobile,
                    "Email": row.Email,
                    "ProgramCode": row.ProgramCode,
                    "CourseName": row.ProgramName,
                    "PreferredBatch": row.PreferredBatch,
                    "StudyMode": row.StudyMode,
                    "AdmissionStatus": row.CourseCompletionStatus,
                    "QrToken": str(row.QrToken)
                    if row.QrToken
                    else None,

                    "Photo": photo_to_data_url(
                        row.PhotoBase64
                    )
                }

                students_data.append(student)

        return render_template(
            "view_admission.html",
            students=students_data,
            search=search
        )

    except Exception as e:

        return (
            f"Unable to load students: {e}",
            500
        )


# ============================================================
# VIEW COMPLETE STUDENT DATA
# ============================================================

@view_student_listing.route(
    "/student/<admission_no>/view",
    methods=["GET"]
)
def view_student(admission_no):

    try:

        with get_sql_server_connection() as connection:

            cursor = connection.cursor()

            sql = """
            SELECT
                AdmissionNo,
                FullName,
                DOB,
                Nationality,
                GovernmentId,
                Gender,
                BloodGroup,
                Category,
                Email,
                Mobile,

                FatherName,
                FatherMobile,
                MotherName,
                MotherMobile,
                GuardianName,
                GuardianMobile,
                FatherOccupation,
                MotherOccupation,
                GuardianIncome,

                PermanentAddress,
                CurrentAddress,
                City,
                State,
                PinCode,

                HighestQualification,
                BoardUniversity,
                SchoolCollege,
                PassingYear,
                PercentageCgpa,
                Medium,

                CourseName,
                PreferredBatch,
                StudyMode,

                DocumentsPath,

                MedicalConditions,
                EmergencyName,
                EmergencyNumber,

                PhotoBase64,

                ApplicantSignature,
                SignatureDate,
                Place,

                ReviewedBy,
                AdmissionStatus,
                QrToken

            FROM tbl_student_admision

            WHERE AdmissionNo = ?
            """

            cursor.execute(
                sql,
                (admission_no,)
            )

            row = cursor.fetchone()

            if not row:
                abort(404)

            student = {
                "AdmissionNo": row.AdmissionNo,
                "FullName": row.FullName,
                "DOB": row.DOB,
                "Nationality": row.Nationality,
                "GovernmentId": row.GovernmentId,
                "Gender": row.Gender,
                "BloodGroup": row.BloodGroup,
                "Category": row.Category,
                "Email": row.Email,
                "Mobile": row.Mobile,

                "FatherName": row.FatherName,
                "FatherMobile": row.FatherMobile,
                "MotherName": row.MotherName,
                "MotherMobile": row.MotherMobile,
                "GuardianName": row.GuardianName,
                "GuardianMobile": row.GuardianMobile,
                "FatherOccupation": row.FatherOccupation,
                "MotherOccupation": row.MotherOccupation,
                "GuardianIncome": row.GuardianIncome,

                "PermanentAddress": row.PermanentAddress,
                "CurrentAddress": row.CurrentAddress,
                "City": row.City,
                "State": row.State,
                "PinCode": row.PinCode,

                "HighestQualification": row.HighestQualification,
                "BoardUniversity": row.BoardUniversity,
                "SchoolCollege": row.SchoolCollege,
                "PassingYear": row.PassingYear,
                "PercentageCgpa": row.PercentageCgpa,
                "Medium": row.Medium,

                "CourseName": row.CourseName,
                "PreferredBatch": row.PreferredBatch,
                "StudyMode": row.StudyMode,

                "DocumentsPath": row.DocumentsPath,

                "MedicalConditions": row.MedicalConditions,
                "EmergencyName": row.EmergencyName,
                "EmergencyNumber": row.EmergencyNumber,

                "Photo": photo_to_data_url(
                    row.PhotoBase64
                ),

                "ApplicantSignature": row.ApplicantSignature,
                "SignatureDate": row.SignatureDate,
                "Place": row.Place,

                "ReviewedBy": row.ReviewedBy,
                "AdmissionStatus": row.AdmissionStatus,

                "QrToken": str(row.QrToken)
                if row.QrToken
                else None
            }

        return render_template(
            "student_view.html",
            student=student
        )

    except Exception as e:

        return (
            f"Unable to load student: {e}",
            500
        )


# ============================================================
# PRINT COMPLETE ADMISSION FORM
# ============================================================

@view_student_listing.route(
    "/student/<admission_no>/print",
    methods=["GET"]
)
def print_student(admission_no):

    try:

        with get_sql_server_connection() as connection:

            cursor = connection.cursor()

            sql = """
            SELECT
                AdmissionNo,
                FullName,
                DOB,
                Nationality,
                GovernmentId,
                Gender,
                BloodGroup,
                Category,
                Email,
                Mobile,

                FatherName,
                FatherMobile,
                MotherName,
                MotherMobile,
                GuardianName,
                GuardianMobile,
                FatherOccupation,
                MotherOccupation,
                GuardianIncome,

                PermanentAddress,
                CurrentAddress,
                City,
                State,
                PinCode,

                HighestQualification,
                BoardUniversity,
                SchoolCollege,
                PassingYear,
                PercentageCgpa,
                Medium,

                CourseName,
                PreferredBatch,
                StudyMode,

                DocumentsPath,

                MedicalConditions,
                EmergencyName,
                EmergencyNumber,

                PhotoBase64,

                ApplicantSignature,
                SignatureDate,
                Place,

                ReviewedBy,
                AdmissionStatus,
                QrToken

            FROM Students

            WHERE AdmissionNo = ?
            """

            cursor.execute(
                sql,
                (admission_no,)
            )

            row = cursor.fetchone()

            if not row:
                abort(404)

            student = {
                "AdmissionNo": row.AdmissionNo,
                "FullName": row.FullName,
                "DOB": row.DOB,
                "Nationality": row.Nationality,
                "GovernmentId": row.GovernmentId,
                "Gender": row.Gender,
                "BloodGroup": row.BloodGroup,
                "Category": row.Category,
                "Email": row.Email,
                "Mobile": row.Mobile,

                "FatherName": row.FatherName,
                "FatherMobile": row.FatherMobile,
                "MotherName": row.MotherName,
                "MotherMobile": row.MotherMobile,
                "GuardianName": row.GuardianName,
                "GuardianMobile": row.GuardianMobile,
                "FatherOccupation": row.FatherOccupation,
                "MotherOccupation": row.MotherOccupation,
                "GuardianIncome": row.GuardianIncome,

                "PermanentAddress": row.PermanentAddress,
                "CurrentAddress": row.CurrentAddress,
                "City": row.City,
                "State": row.State,
                "PinCode": row.PinCode,

                "HighestQualification": row.HighestQualification,
                "BoardUniversity": row.BoardUniversity,
                "SchoolCollege": row.SchoolCollege,
                "PassingYear": row.PassingYear,
                "PercentageCgpa": row.PercentageCgpa,
                "Medium": row.Medium,

                "CourseName": row.CourseName,
                "PreferredBatch": row.PreferredBatch,
                "StudyMode": row.StudyMode,

                "DocumentsPath": row.DocumentsPath,

                "MedicalConditions": row.MedicalConditions,
                "EmergencyName": row.EmergencyName,
                "EmergencyNumber": row.EmergencyNumber,

                "Photo": photo_to_data_url(
                    row.PhotoBase64
                ),

                "ApplicantSignature": row.ApplicantSignature,
                "SignatureDate": row.SignatureDate,
                "Place": row.Place,

                "ReviewedBy": row.ReviewedBy,
                "AdmissionStatus": row.AdmissionStatus,

                "QrToken": str(row.QrToken)
                if row.QrToken
                else None
            }

        # ----------------------------------------------------
        # Generate QR
        # ----------------------------------------------------

        qr_token = student["QrToken"]

        qr_image = generate_qr_image(
            qr_token
        )

        qr_base64 = None

        if qr_image:

            qr_encoded = base64.b64encode(
                qr_image.getvalue()
            ).decode("utf-8")

            qr_base64 = (
                "data:image/png;base64,"
                + qr_encoded
            )
  
        return render_template(
            "student_print.html",
            student=student,
            qr_code=qr_base64
        )

    except Exception as e:

        return (
            f"Unable to print student admission: {e}",
            500
        )


# ============================================================
# OPTIONAL: QR IMAGE ENDPOINT
# ============================================================

@view_student_listing.route(
    "/student/<admission_no>/qr",
    methods=["GET"]
)
def student_qr(admission_no):

    try:

        with get_sql_server_connection() as connection:

            cursor = connection.cursor()

            sql = """
            SELECT QrToken
            FROM Students
            WHERE AdmissionNo = ?
            """

            cursor.execute(
                sql,
                (admission_no,)
            )

            row = cursor.fetchone()

            if not row or not row.QrToken:
                abort(404)

            qr_token = str(
                row.QrToken
            )

        qr_image = generate_qr_image(
            qr_token
        )

        if not qr_image:
            abort(404)

        return send_file(
            qr_image,
            mimetype="image/png",
            download_name=f"{admission_no}.png"
        )

    except Exception as e:

        return (
            f"Unable to generate QR: {e}",
            500
        )

@view_student_listing.route("/view_students/export-excel", methods=["GET"])
def export_students_excel():

    search = request.args.get("search", "").strip()

    connection = None
    cursor = None

    try:
        connection = get_sql_server_connection()
        cursor = connection.cursor()

        if search:

            sql = """
                SELECT
                    AdmissionNo,
                    FullName,
                    DOB,
                    Gender,
                    Mobile,
                    Email,
                    ProgramName,
                    PreferredBatch,
                    StudyMode,
                    CourseCompletionStatus
                FROM tbl_student_admision
                WHERE
                    AdmissionNo LIKE ?
                    OR FullName LIKE ?
                    OR Mobile LIKE ?
                    OR Email LIKE ?
                ORDER BY AdmissionNo DESC
            """

            search_value = f"%{search}%"

            cursor.execute(
                sql,
                (
                    search_value,
                    search_value,
                    search_value,
                    search_value
                )
            )

        else:

            sql = """
                SELECT
                    ID,
                    AdmissionNo,
                    FullName,
                    DOB,
                    Gender,
                    Mobile,
                    Email,
                    ProgramName,
                    PreferredBatch,
                    StudyMode,
                    CourseCompletionStatus,
                    QrToken
                FROM tbl_student_admision
                WHERE CourseCompletionStatus NOT IN
                    ('Hold', 'Completed', 'DROPOUT')
                ORDER BY AdmissionNo DESC
            """

            cursor.execute(sql)

        rows = cursor.fetchall()

        # --------------------------------------------------
        # Create Excel workbook
        # --------------------------------------------------

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Student Admissions"

        headers = [
            "Enrollment No",
            "Admission No",
            "Full Name",
            "DOB",
            "Gender",
            "Mobile",
            "Email",
            "Course",
            "Preferred Batch",
            "Study Mode",
            "Status",
            "QR Token"
        ]

        worksheet.append(headers)

        # Header formatting
        for cell in worksheet[1]:

            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )

            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="1F6FEB"
            )

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

        # --------------------------------------------------
        # Add student data
        # --------------------------------------------------

        for row in rows:

            worksheet.append([
                row.ID,
                row.AdmissionNo,
                row.FullName,
                row.DOB,
                row.Gender,
                row.Mobile,
                row.Email,
                row.ProgramName,
                row.PreferredBatch,
                row.StudyMode,
                row.CourseCompletionStatus,
                str(row.QrToken) if row.QrToken else ""

            ])

        # --------------------------------------------------
        # Column widths
        # --------------------------------------------------

        column_widths = {
            "A": 18,
            "B": 30,
            "C": 15,
            "D": 12,
            "E": 18,
            "F": 30,
            "G": 25,
            "H": 20,
            "I": 18,
            "J": 20
        }

        for column, width in column_widths.items():
            worksheet.column_dimensions[column].width = width

        # Freeze header
        worksheet.freeze_panes = "A2"

        # Auto filter
        worksheet.auto_filter.ref = worksheet.dimensions

        # --------------------------------------------------
        # Save to memory
        # --------------------------------------------------

        output = BytesIO()

        workbook.save(output)

        output.seek(0)

        return send_file(
            output,
            as_attachment=True,
            download_name="student_admissions.xlsx",
            mimetype=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )
        )

    except Exception as e:

        print("Excel export error:", e)

        return (
            f"Error exporting students to Excel: {str(e)}",
            500
        )

    finally:

        try:
            if cursor:
                cursor.close()
        except Exception:
            pass

        try:
            if connection:
                connection.close()
        except Exception:
            pass

